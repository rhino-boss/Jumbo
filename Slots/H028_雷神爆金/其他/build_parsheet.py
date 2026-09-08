# -*- coding: utf-8 -*-
"""Generate H028 PARsheet workbooks (one per RTP variant), modeled on H021 PARsheet.
Content: Overview (model/version, RTP breakdown per profile, paytable per 100 credits,
FG spins, M1 mapping, caps, SCR) + Parameter (table-selection weights) + 6 symbol sheets
(values-only copies from H0281.xlsx). Card internals (Multiplier_Weight/Detail) excluded."""
import sys, io, json, re, os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
BASE = r"C:\Users\rhinshen\Mine\個人工作區\工作區\Slots\H028_雷神爆金"
OUT = os.path.join(BASE, "PARsheet")
os.makedirs(OUT, exist_ok=True)

TAGS = ['88B', '90B', '92A', '94A']
FG_TARGET = {'88B': 0.16, '90B': 0.18, '92A': 0.20, '94A': 0.22}
SYMBOL_SHEETS = ['BG_Symbol', 'BG_Symbol (2)', 'BF_Symbol', 'FG_Symbol', 'FG_Symbol (2)', 'FG_Symbol (3)']

bold = Font(bold=True)
title_font = Font(bold=True, size=12)
hdr_fill = PatternFill('solid', fgColor='DDE6F2')
thin = Side(style='thin', color='B9B9B9')
box = Border(left=thin, right=thin, top=thin, bottom=thin)


def cfg_version(tag):
    txt = open(rf"{BASE}\config_{tag}.js", encoding='utf-8').read()
    return re.search(r'"excel_version":\s*"([^"]+)"', txt).group(1)


def cfg_arrays(tag):
    txt = open(rf"{BASE}\config_{tag}.js", encoding='utf-8').read()
    def arr(owner, sub, key):
        o = txt.find(f'"{owner}"'); s = txt.find(f'"{sub}"', o); a = txt.find(f'"{key}"', s)
        j = txt.find('[', a); d = 0
        for k in range(j, len(txt)):
            if txt[k] == '[': d += 1
            elif txt[k] == ']':
                d -= 1
                if d == 0: return json.loads(txt[j:k + 1])
    def cap(cards):
        vals = [c['max'] for c in cards if c.get('type') == 'range' and int(c.get('weight', 0)) > 0]
        return int(max(vals)) if vals else None
    def fgT(cards):
        return sum(int(c['weight']) for c in cards if c.get('type') == 'free_game')
    o_bg = arr('oldhand', 'normal_bet', 'weight_bg'); o_fg = arr('oldhand', 'normal_bet', 'weight_fg')
    n_bg = arr('newbie', 'normal_bet', 'weight_bg'); n_fg = arr('newbie', 'normal_bet', 'weight_fg')
    bf = arr('oldhand', 'buy_feature', 'weight_fg')
    return {
        'cap_o_bg': cap(o_bg), 'cap_o_fg': cap(o_fg),
        'cap_n_bg': cap(n_bg), 'cap_n_fg': cap(n_fg), 'cap_bf': cap(bf),
        'cycle_o': 1e9 / fgT(o_bg), 'cycle_n': 1e9 / fgT(n_bg),
    }


def read_scr(tag):
    wb = openpyxl.load_workbook(rf"{BASE}\Source\H0281{tag}.xlsx", read_only=True, data_only=True)
    ws = wb['OP Jackpot']
    vals = {}
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None]
        if len(cells) >= 2 and str(cells[0]) in ('NB_Newbie', 'NB', 'BF', 'Threshold'):
            vals[str(cells[0])] = cells[1]
    wb.close()
    return vals


# ---- load base model data once ----
base_wb = openpyxl.load_workbook(rf"{BASE}\Source\H0281.xlsx", data_only=True)
pay_rows = []           # (symbol, description, p3, p4, p5, p6)
fs_rows = []            # (c1_num, spins)
window = None
in_pay = in_fs = False
for row in base_wb['Overview'].iter_rows(values_only=True):
    vals = list(row)
    a = str(vals[0]) if vals[0] is not None else ''
    if a == 'Visible Window Size':
        window = [v for v in vals[1:8] if v is not None]
    if a.startswith('Pay Table'):
        in_pay = True; in_fs = False; continue
    if a == 'Free Spins Setting':
        in_fs = True; continue
    if in_fs:
        if a in ('C1 Num',):
            continue
        if a and (a.isdigit() or a == '…'):
            fs_rows.append((vals[0], vals[1]))
            continue
        if a.startswith('The maximum'):
            in_fs = False
            continue
    if in_pay:
        if a == 'Symbol':
            continue
        if a and vals[6] is not None:  # has Id column
            pay_rows.append(tuple(vals[0:2]) + tuple(vals[2:6]))
        elif not a:
            in_pay = False

param_rows = [[c for c in row] for row in base_wb['Parameter'].iter_rows(values_only=True)]


def put(ws, r, c, v, font=None, fill=None, border=True):
    cell = ws.cell(row=r, column=c, value=v)
    if font: cell.font = font
    if fill: cell.fill = fill
    if border: cell.border = box
    return cell


def build_overview(ws, tag, caps, scr):
    fg = FG_TARGET[tag]
    total = round(0.72 + fg, 4)
    ws.column_dimensions['A'].width = 26
    for col in 'BCDEFG':
        ws.column_dimensions[col].width = 15
    r = 1
    put(ws, r, 1, 'Model:', bold, border=False); put(ws, r, 2, f'H0281{tag}', border=False); r += 1
    put(ws, r, 1, 'Version:', bold, border=False); put(ws, r, 2, cfg_version(tag), border=False); r += 1
    put(ws, r, 1, 'Hold:', bold, border=False); put(ws, r, 2, round(1 - total, 6), border=False); r += 2

    put(ws, r, 1, 'Base Bet', bold, hdr_fill); put(ws, r, 2, 'Min Ways', bold, hdr_fill); put(ws, r, 3, 'Max Ways', bold, hdr_fill); r += 1
    put(ws, r, 1, 100); put(ws, r, 2, 2025); put(ws, r, 3, 32400); r += 2

    hdr = ['Bet Type', 'Coin in', 'Price(x)', 'Profile', 'Base Game Pay Back', 'Free Game Pay Back', 'Total RTP']
    for i, h in enumerate(hdr, 1): put(ws, r, i, h, bold, hdr_fill)
    r += 1
    rows = [
        ('Normal Bet', 100, 1, 'Oldhand', 0.72, fg, total),
        ('Normal Bet', 100, 1, 'Newbie', 0.72, 0.21, 0.93),
        ('Buy Feature', 7500, 75, 'Oldhand / Newbie', 0.0, 0.925, 0.925),
    ]
    for row_vals in rows:
        for i, v in enumerate(row_vals, 1): put(ws, r, i, v)
        r += 1
    r += 1

    hdr = ['Profile', 'Free Game Hits', 'Pulls/Hit (FG Cycle)']
    for i, h in enumerate(hdr, 1): put(ws, r, i, h, bold, hdr_fill)
    r += 1
    put(ws, r, 1, 'Oldhand Normal Bet'); put(ws, r, 2, round(1 / caps['cycle_o'], 10)); put(ws, r, 3, round(caps['cycle_o'], 5)); r += 1
    put(ws, r, 1, 'Newbie Normal Bet'); put(ws, r, 2, round(1 / caps['cycle_n'], 10)); put(ws, r, 3, round(caps['cycle_n'], 5)); r += 1
    put(ws, r, 1, 'Buy Feature'); put(ws, r, 2, 1); put(ws, r, 3, 1); r += 2

    put(ws, r, 1, 'Scatter symbol appear rate (SCR, x1e10)', bold, hdr_fill)
    put(ws, r, 2, 'NB Oldhand', bold, hdr_fill); put(ws, r, 3, 'NB Newbie', bold, hdr_fill); put(ws, r, 4, 'Buy Feature', bold, hdr_fill); r += 1
    put(ws, r, 1, '')
    put(ws, r, 2, scr.get('NB')); put(ws, r, 3, scr.get('NB_Newbie')); put(ws, r, 4, scr.get('BF')); r += 2

    put(ws, r, 1, 'Reel #', bold, hdr_fill)
    for i in range(6): put(ws, r, 2 + i, i + 1, bold, hdr_fill)
    r += 1
    put(ws, r, 1, 'Visible Window Size (incl. Extra Reel on R2-R5)')
    for i, v in enumerate((window or [5, 6, 6, 6, 6, 5])[:6]): put(ws, r, 2 + i, v)
    r += 2

    put(ws, r, 1, 'Free Spins Setting', title_font, border=False); r += 1
    put(ws, r, 1, 'Scatter Num', bold, hdr_fill); put(ws, r, 2, 'Free Spins', bold, hdr_fill); r += 1
    for c1, spins in fs_rows:
        put(ws, r, 1, c1); put(ws, r, 2, spins); r += 1
    put(ws, r, 1, 'Each additional Scatter awards +2 spins; maximum 50 free spins per feature. Scatters on main reels and extra reels are all counted.', border=False); r += 2

    put(ws, r, 1, 'Pay Table:', title_font, border=False)
    put(ws, r, 2, 'All wins show for 100 credit bets', border=False); r += 1
    hdr = ['Symbol', 'Description', '3', '4', '5', '6']
    for i, h in enumerate(hdr, 1): put(ws, r, i, h, bold, hdr_fill)
    r += 1
    for row_vals in pay_rows:
        for i, v in enumerate(row_vals, 1): put(ws, r, i, v)
        r += 1
    put(ws, r, 1, 'WW substitutes for all symbols except C1 (Scatter). C1 / M1 golden-frame symbols pay the same as base symbols. A large symbol counts as 1 symbol per reel for Way calculation.', border=False); r += 2

    put(ws, r, 1, 'M1 Multiplier (main reels, by symbol size)', title_font, border=False); r += 1
    put(ws, r, 1, 'Size', bold, hdr_fill); put(ws, r, 2, 'Multiplier', bold, hdr_fill); r += 1
    for size, mult in (('1x1', 'x2'), ('1x2', 'x3'), ('1x3', 'x4'), ('1x4', 'x5')):
        put(ws, r, 1, size); put(ws, r, 2, mult); r += 1
    put(ws, r, 1, 'Each M1 on the Extra Reels awards a fixed x2. Multipliers in a round are added together. Free Game starts at x2 and is carried over across free spins.', border=False); r += 2

    put(ws, r, 1, 'Feature / Limits', title_font, border=False); r += 1
    limits = [
        ('Buy Feature price', '75x total bet (direct Free Game entry)'),
        ('Card multiplier cap - Oldhand BG / FG', f"{caps['cap_o_bg']}x / {caps['cap_o_fg']:,}x"),
        ('Card multiplier cap - Newbie BG / FG', f"{caps['cap_n_bg']}x / {caps['cap_n_fg']}x"),
        ('Card multiplier cap - Buy Feature FG', f"{caps['cap_bf']:,}x"),
        ('Golden frame', 'General symbols on R2-R5 may be gold-framed; turns into WW after participating in a win (kept for 1 cascade).'),
        ('Jackpot', 'OP Jackpot (platform feature): GRAND / MAJOR (linked progressive, unlocked at bet 2.00+), MINOR / MINI (bonus).'),
    ]
    for k, v in limits:
        put(ws, r, 1, k); put(ws, r, 2, v, border=False); r += 1


def build_parameter(ws):
    labels = ['Base Game Table Selection', 'Free Game Initial Table Selection', 'Free Game Retrigger Table Selection']
    li = -1
    r_out = 1
    ws.column_dimensions['B'].width = 18
    ws.column_dimensions['C'].width = 12
    for row in param_rows:
        vals = row
        if vals[1] == 'Table Selection Weight':
            li += 1
            put(ws, r_out, 1, labels[li], title_font, border=False); r_out += 1
            continue
        if vals[1] is not None:
            put(ws, r_out, 2, vals[1], bold if vals[1] == 'Worksheet Name' else None,
                hdr_fill if vals[1] == 'Worksheet Name' else None)
            put(ws, r_out, 3, vals[2], bold if vals[2] == 'Weight' else None,
                hdr_fill if vals[2] == 'Weight' else None)
            r_out += 1
    r_out += 1
    put(ws, r_out, 1, 'Free Game spins: 4 Scatters award 10 spins, each additional Scatter +2, maximum 50 per feature.', border=False)


def copy_symbol_sheet(dst_ws, src_ws):
    for row in src_ws.iter_rows():
        for cell in row:
            if cell.value is not None:
                dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)


for tag in TAGS:
    caps = cfg_arrays(tag)
    scr = read_scr(tag)
    wb = openpyxl.Workbook()
    ov = wb.active
    ov.title = 'Overview'
    build_overview(ov, tag, caps, scr)
    build_parameter(wb.create_sheet('Parameter'))
    for name in SYMBOL_SHEETS:
        copy_symbol_sheet(wb.create_sheet(name), base_wb[name])
    out = os.path.join(OUT, f'H0281{tag}.xlsx')
    wb.save(out)
    print('written', out)
print('done')
