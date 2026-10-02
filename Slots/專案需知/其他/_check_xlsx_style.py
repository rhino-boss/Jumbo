# -*- coding: utf-8 -*-
"""xlsx 格式一致性檢查器。

用途：程式化修改數學模型 xlsx 後，檢查有沒有「空白格殘留格式」與「有值卻漏格式」。
起因：openpyxl 把 `.value` 設成 None 只清值不清格式；新寫入的儲存格也不會自動繼承格式。
兩者疊起來就會出現空白處有底色外框、有值處卻是裸格。

用法：
    py _check_xlsx_style.py <xlsx 路徑> [工作表名稱 ...]

判定：
    [ERROR] 空白格帶有底色                  —— 幾乎必為殘留，必須清掉
    [ERROR] 空白格有框線且同列完全沒有值      —— 孤立的殘留表格
    [WARN ] 同一列其他格有格式，本格有值卻全裸 —— 多半是漏套，需人工確認
    [WARN ] 有值格的字體與該表主要字體不同    —— 多半是新寫入沒複製樣式

    空白格只有框線、且同列有值者視為正常（表格補格，避免表格缺角）。

離開碼：有 ERROR 回傳 1，其餘回傳 0。
"""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

import openpyxl


def has_fill(cell) -> bool:
    fill = cell.fill
    if fill is None or fill.patternType is None:
        return False
    return getattr(fill.fgColor, "rgb", None) not in (None, "00000000")


def has_border(cell) -> bool:
    border = cell.border
    return any(getattr(border, side).style for side in ("left", "right", "top", "bottom"))


def styled(cell) -> bool:
    return has_fill(cell) or has_border(cell)


def check_sheet(ws) -> tuple[list[str], list[str]]:
    errors: list[str] = []
    warnings: list[str] = []

    fonts = Counter()
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is not None and cell.font.name:
                fonts[cell.font.name] += 1
    main_font = fonts.most_common(1)[0][0] if fonts else None

    for row in ws.iter_rows():
        cells = [c for c in row if c.value is not None or styled(c)]
        if not cells:
            continue
        row_has_style = any(styled(c) for c in cells)
        row_has_value = any(c.value is not None for c in row)
        for cell in row:
            value = cell.value is not None
            if not value and has_fill(cell):
                errors.append(f"{ws.title}!{cell.coordinate} 空白卻有底色")
            elif not value and has_border(cell) and not row_has_value:
                errors.append(f"{ws.title}!{cell.coordinate} 空白且同列無值卻有外框")
            elif value and not styled(cell) and row_has_style:
                warnings.append(f"{ws.title}!{cell.coordinate} 有值卻無底色無外框"
                                f"（同列其他格有格式）")
            if value and main_font and cell.font.name and cell.font.name != main_font:
                warnings.append(f"{ws.title}!{cell.coordinate} 字體 {cell.font.name} "
                                f"≠ 本表主要字體 {main_font}")
    return errors, warnings


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"找不到檔案：{path}")
        return 2
    wanted = sys.argv[2:]

    wb = openpyxl.load_workbook(path)
    sheets = [wb[name] for name in wanted] if wanted else wb.worksheets

    all_errors: list[str] = []
    all_warnings: list[str] = []
    for ws in sheets:
        errors, warnings = check_sheet(ws)
        all_errors += errors
        all_warnings += warnings

    print(f"檔案：{path.name}")
    print(f"檢查工作表：{', '.join(ws.title for ws in sheets)}")
    print()
    for item in all_errors:
        print(f"[ERROR] {item}")
    for item in all_warnings:
        print(f"[WARN ] {item}")
    print()
    print(f"ERROR {len(all_errors)} 筆、WARN {len(all_warnings)} 筆")
    if not all_errors and not all_warnings:
        print("格式一致，無殘留也無漏套。")
    elif not all_errors:
        print("無殘留；WARN 請人工確認（區塊標題本來就無框屬正常）。")
    return 1 if all_errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
